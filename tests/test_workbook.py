"""Workbook integrity without LibreOffice: data cells sum to the source total and every
=SUM() totals formula spans exactly the data rows. Runs on the fixture and, if present,
on the real docs/aero-supplier-spend.xlsx against data/spend.duckdb."""
import re
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from src.crosswalk import build_supplier, load_crosswalk
from src.export_xlsx import build
from src.load import build_db, zip_to_parquet
from src.metrics import run
from tests.helpers import txn, write_zip

SUM_RE = re.compile(r"^=SUM\(([A-Z]+)2:([A-Z]+)(\d+)\)$")


def check_workbook(path, expected_total):
    """Return a list of problems; empty means the workbook is consistent."""
    problems = []
    wb = load_workbook(path)
    ws = wb["Supplier Cube"]
    header = [c.value for c in ws[1]]
    col = header.index("Obligations ($)") + 1
    total_row = ws.max_row
    data = [ws.cell(r, col).value for r in range(2, total_row)]
    if round(Decimal(str(sum(data))), 2) != round(Decimal(str(expected_total)), 2):
        problems.append(f"cube data sum {sum(data)} != source {expected_total}")
    for sheet in wb.worksheets:
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("=SUM"):
                m = SUM_RE.match(cell.value)
                if not m or int(m.group(3)) != sheet.max_row - 1:
                    problems.append(f"{sheet.title}!{cell.coordinate} range wrong: {cell.value}")
    return problems


@pytest.fixture
def fixture_wb(tmp_path):
    write_zip(tmp_path / "dod-fy2025.zip", [txn("T1", 50), txn("T2", 30, recipient_uei="U2"), txn("T3", -5)])
    zip_to_parquet(tmp_path / "dod-fy2025.zip", tmp_path)
    con = build_db(str(tmp_path / "*.parquet"), tmp_path / "t.duckdb")
    load_crosswalk(con, tmp_path / "none.csv")
    build_supplier(con)
    run(con, tmp_path / "out")
    return build(con, tmp_path / "w.xlsx", tmp_path / "out", tmp_path / "none.csv")


def test_fixture_workbook_consistent(fixture_wb):
    assert check_workbook(fixture_wb, 75) == []


def test_dropped_row_goes_red(fixture_wb):
    wb = load_workbook(fixture_wb)
    wb["Supplier Cube"].delete_rows(2)
    wb.save(fixture_wb)
    assert check_workbook(fixture_wb, 75)


@pytest.mark.skipif(not Path("docs/aero-supplier-spend.xlsx").exists(), reason="real workbook not built")
def test_real_workbook_matches_db():
    con = duckdb.connect("data/spend.duckdb", read_only=True)
    total = con.execute("SELECT sum(obligation) FROM fact_sup").fetchone()[0]
    assert check_workbook("docs/aero-supplier-spend.xlsx", total) == []


def evaluated_totals(path):
    """Evaluate every totals-row formula with a Python Excel engine (no LibreOffice needed);
    return (sheet, cell, engine value, data-cell sum) for each."""
    import os

    import formulas
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    by_key = {k.upper(): v for k, v in sol.items()}
    out = []
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                key = f"'[{os.path.basename(path)}]{ws.title}'!{cell.coordinate}".upper()
                val = by_key[key].value[0][0]
                data = sum((ws.cell(r, cell.column).value or 0) for r in range(2, ws.max_row))
                out.append((ws.title, cell.coordinate, val, data))
    return out


def test_formulas_evaluate_to_data_sums(fixture_wb):
    rows = evaluated_totals(fixture_wb)
    assert rows and all(abs(v - d) < 0.01 for _, _, v, d in rows)


@pytest.mark.skipif(not Path("docs/aero-supplier-spend.xlsx").exists(), reason="real workbook not built")
def test_real_formulas_evaluate_to_data_sums():
    rows = evaluated_totals(Path("docs/aero-supplier-spend.xlsx"))
    assert rows and all(abs(v - d) < 0.01 for _, _, v, d in rows)
